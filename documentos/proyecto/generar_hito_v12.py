#!/usr/bin/env python3
"""Publica solo artefactos v1.2 desde fuentes Markdown; no ejecuta la PoC.

Requiere bibliotecas existentes. DOCX y PDF se generan independientemente;
no se afirma equivalencia de paginación ni renderizado de Word.
"""

from html import escape
from math import ceil
from pathlib import Path
import re
import sys

from artefactos_hito import bloques, leer_tablas, sin_marcas

ROOT = Path(__file__).resolve().parents[2]
FUENTE = ROOT / "documentos/proyecto/Semana5_Hito_I_v1.2.md"
DESTINO = ROOT / "entregas/semana-05/revision-v1.2"
NOMBRE = "TISG_E01_S05_HitoI_v1.2_2026-10-25"
LIBRO = ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.2.xlsx"


def docx(texto, fuente, salida, titulo):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.7)
    s.left_margin = s.right_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for tipo, valor in bloques(texto):
        if tipo == "titulo":
            nivel, contenido = valor
            doc.add_heading(sin_marcas(contenido), level=min(nivel, 3))
        elif tipo == "tabla":
            t = doc.add_table(rows=1, cols=len(valor[0]))
            t.style = "Light Shading Accent 1"
            for c, contenido in zip(t.rows[0].cells, valor[0]):
                c.text = sin_marcas(contenido)
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for fila in valor[1:]:
                for c, contenido in zip(t.add_row().cells, fila):
                    c.text = sin_marcas(contenido)
            for fila in t.rows:
                fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for c in fila.cells:
                    for p in c.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8)
            doc.add_paragraph()
        elif tipo == "imagen":
            alt, ruta = valor
            doc.add_picture(str((fuente.parent / ruta).resolve()), width=Cm(17.7))
            doc.add_paragraph(alt, style="Caption")
        else:
            p = doc.add_paragraph()
            for fragmento in re.split(r"(\*\*.*?\*\*|`.*?`)", valor):
                run = p.add_run(sin_marcas(fragmento))
                run.bold = fragmento.startswith("**")
    pie = s.footer.paragraphs[0]
    pie.text = "Grupo 1 · v1.2 · 25/10/2026 · Pendiente de validación humana · "
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)
    doc.core_properties.title = titulo
    doc.core_properties.author = "Grupo 1; revisión humana pendiente"
    doc.save(salida)


def pdf(texto, salida):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Image, Spacer, KeepTogether
    from PIL import Image as PILImage

    fonts = Path("C:/Windows/Fonts")
    for nombre, archivo in [("Hito12", "arial.ttf"), ("Hito12-Bold", "arialbd.ttf")]:
        pdfmetrics.registerFont(TTFont(nombre, str(fonts / archivo)))
    pdfmetrics.registerFontFamily("Hito12", normal="Hito12", bold="Hito12-Bold", italic="Hito12", boldItalic="Hito12-Bold")
    estilos = getSampleStyleSheet()
    for estilo in estilos.byName.values():
        estilo.fontName = "Hito12"
    estilos["BodyText"].fontSize, estilos["BodyText"].leading = 9, 13
    for nombre in ("Heading1", "Heading2", "Heading3"):
        estilos[nombre].keepWithNext = True
    estilos.add(ParagraphStyle("Celda12", fontName="Hito12", fontSize=7.4, leading=10, splitLongWords=True))
    estilos.add(ParagraphStyle("Pie12", fontName="Hito12", fontSize=8, leading=10, spaceAfter=10))
    ancho = A4[0] - 3.2 * cm

    def xml(t):
        return re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escape(t).replace("`", ""))

    elementos = []
    for tipo, valor in bloques(texto):
        if tipo == "titulo":
            nivel, titulo = valor
            nombre = "Title" if nivel == 1 else "Heading" + str(min(nivel - 1, 3))
            elementos.append(Paragraph(xml(titulo), estilos[nombre]))
        elif tipo == "tabla":
            datos = [[Paragraph(xml(c), estilos["Celda12"]) for c in f] for f in valor]
            pesos = [max(14, min(60, max(len(sin_marcas(f[j])) for f in valor))) for j in range(len(valor[0]))]
            minimo = 30
            anchos = [minimo + (ancho - minimo * len(pesos)) * p / sum(pesos) for p in pesos]
            tabla = LongTable(datos, colWidths=anchos, repeatRows=1, hAlign="LEFT")
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            elementos.extend([tabla, Spacer(1, 9)])
        elif tipo == "imagen":
            alt, relativa = valor
            ruta = (FUENTE.parent / relativa).resolve()
            with PILImage.open(ruta) as im:
                w, h = im.size
            elementos.append(KeepTogether([Image(str(ruta), width=ancho, height=ancho * h / w), Paragraph(xml(alt), estilos["Pie12"])]))
        else:
            elementos.append(Paragraph(xml(valor), estilos["BodyText"]))

    def pie(canvas, documento):
        canvas.setFont("Hito12", 7.5)
        canvas.drawString(1.6 * cm, 1.0 * cm, "Grupo 1 · Hito I v1.2 · 25/10/2026 · Pendiente de validación humana")
        canvas.drawRightString(A4[0] - 1.6 * cm, 1.0 * cm, str(documento.page))

    SimpleDocTemplate(str(salida), pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                      topMargin=1.6 * cm, bottomMargin=1.7 * cm,
                      title="Hito I v1.2 — Grupo 1", author="Grupo 1; validación humana pendiente").build(elementos, onFirstPage=pie, onLaterPages=pie)


def libro(texto):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    anterior = ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.1.xlsx"
    w = load_workbook(anterior, data_only=False, keep_links=True)
    for nombre, nuevo in [("Madurez_SocioTecnica", "Madurez_Hist_v11"), ("Fuentes", "Fuentes_Hist_v11"), ("Hipotesis", "Hipotesis_Hist_v11"), ("Supuestos", "Supuestos_Hist_v11")]:
        copia = w.copy_worksheet(w[nombre])
        copia.title = nuevo
    tablas = leer_tablas(texto)

    def tabla(cabecera):
        coincidencias = [t for t in tablas if t[0][0] == cabecera]
        if len(coincidencias) != 1:
            raise ValueError(f"Tabla no única: {cabecera}")
        return coincidencias[0]

    def escribir(nombre, filas):
        s = w[nombre] if nombre in w.sheetnames else w.create_sheet(nombre)
        s.delete_rows(1, s.max_row)
        for fila in filas:
            s.append([sin_marcas(str(v)) if v is not None else None for v in fila])

    escribir("Control", [["Campo", "Valor"], ["Versión", "v1.2; 25/10/2026; S5 integra S1–S5"],
                          ["Estado", "Pendiente de validación humana"], ["Fuente", "documentos/proyecto/Semana5_Hito_I_v1.2.md"],
                          ["Conservación", "Copiado de v1.1; históricos y contribuciones conservados"],
                          ["Madurez", "Instrumento propio; ND no es cero; solo datos 2-H hipotético"],
                          ["Experimento", "Protocolo nuevo no ejecutado; resultados previos solo históricos"],
                          ["Fórmulas", "El libro de registros v1.1 no contenía fórmulas; definiciones metodológicas en texto, sin recálculo simulado"],
                          ["Confidencialidad", "Caso ficticio; no conexión productiva ni inventario real"],
                          ["Formato", "Integridad y contenido comprobables; diseño Excel pendiente de revisión en escritorio"]])
    for nombre, cabecera in {
        "ADR": "ID", "Radar_Data": "ID", "Canvas_Data": "Bloque BMC AS-IS",
        "Propuestas": "Propuesta", "Stakeholders": "Stakeholder",
        "Madurez_SocioTecnica": "Dimensión evaluada", "Skills Gap": "Rol del escenario",
        "Impact_Canvas": "Tendencia", "Hipotesis": "Hipótesis", "Deuda_Trazabilidad": "ID",
        "Dependencias": "ID", "Instrumento_Madurez": "Dimensión del instrumento", "Metricas": "Métrica",
        "Umbrales": "Umbral interno", "Datos_Relaciones": "Relación conceptual", "Topologia": "Elemento",
        "Revision_Humana": "Control humano", "Coherencia": "Vínculo",
    }.items():
        if cabecera == "ID":
            primero = {"ADR": "ADR-01", "Radar_Data": "T1", "Deuda_Trazabilidad": "DT-01", "Dependencias": "F-01"}[nombre]
            filas = next(t for t in tablas if len(t) > 1 and t[1][0] == primero)
        else:
            filas = tabla(cabecera)
        escribir(nombre, filas)

    decisiones = w["Decisiones"]
    for fila in tabla("ID nuevo")[1:]:
        decisiones.append([fila[0], fila[1], fila[2], "Conservar enfoque previo sin corregir observación", fila[3], "Equipo, aprobación pendiente", fila[4]])
    for fila in decisiones.iter_rows(min_row=2):
        if fila[0].value == "D-10":
            fila[4].value = str(fila[4].value) + "; v1.2: retirar fundamento CAPECE H1 2026, mantener alcance contractual supuesto (D-21)"
        if fila[0].value == "D-12":
            fila[6].value = str(fila[6].value) + "; resultados anteriores históricos; nuevo diseño D-19, sin ejecución"
    for fila in tabla("Supuesto nuevo")[1:]:
        w["Supuestos"].append([fila[0], fila[1], "Escenario v1.2", "Propuesto, no observado", fila[2]])
    for fila in tabla("Cambio v1.2")[1:]:
        w["Cambios"].append(fila)

    fuentes = leer_tablas((ROOT / "fuentes/revision-v1.2/VERIFICACION_FUENTES.md").read_text(encoding="utf-8"))
    escribir("Fuentes", next(t for t in fuentes if t[0][0] == "ID"))
    escribir("CAPECE_Correcciones", next(t for t in fuentes if t[0][0] == "ID de corrección"))
    aportes = leer_tablas((ROOT / "gestion/registros/Contribuciones_Evidencia_v1.2.md").read_text(encoding="utf-8"))
    escribir("Evidencia_Aportes", aportes[0])
    escribir("Matriz_BMC_Laminas", [["Bloque", "Fuente", "Lámina", "Observación"], *[
        [fila[0], "Hito I v1.2 §3.1", "Presentacion_S2_v1.2.pdf p. 1", "Mismo contenido; canales resumidos, comerciales/digitales separados"]
        for fila in tabla("Bloque BMC AS-IS")[1:]]])

    for s in w:
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        for c in s[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor="1F4E79")
        for columna in range(1, s.max_column + 1):
            longitud = max(len(str(s.cell(f, columna).value or "")) for f in range(1, s.max_row + 1))
            s.column_dimensions[get_column_letter(columna)].width = min(58, max(16, longitud * .6))
        for fila in s.iter_rows(min_row=2):
            alto = 35
            for c in fila:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(name="Arial", size=10)
                ancho = s.column_dimensions[c.column_letter].width
                alto = max(alto, 15 * ceil(len(str(c.value or "")) / max(8, int(ancho * .8))) + 12)
            s.row_dimensions[fila[0].row].height = min(400, alto)
    w.properties.title = "Registros del Hito I v1.2 — Grupo 1"
    w.properties.description = "Pendiente de validación humana; diagnóstico hipotético; nuevo experimento no ejecutado"
    w.save(LIBRO)
    w.close()


def generar():
    DESTINO.mkdir(parents=True, exist_ok=True)
    texto = FUENTE.read_text(encoding="utf-8")
    docx(texto, FUENTE, DESTINO / (NOMBRE + ".docx"), "Hito I v1.2 — Grupo 1")
    pdf(texto, DESTINO / (NOMBRE + ".pdf"))
    libro(texto)
    fuente_s2 = ROOT / "documentos/proyecto/Semana2_MaterialComplementario_v1.2.md"
    docx(fuente_s2.read_text(encoding="utf-8"), fuente_s2,
         ROOT / "entregas/semana-02/revision-v1.2/MaterialComplementario_S2_v1.2.docx", "S2 — Material complementario v1.2")
    print("[OK] Hito I DOCX/PDF, registros XLSX y material S2 DOCX generados solo en v1.2")


if __name__ == "__main__":
    generar()
