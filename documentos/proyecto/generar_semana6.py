"""Genera SOLO borradores S6 en destinos nuevos; nunca llama generadores históricos.

Escribe arquitectura/semana-06, entregas/semana-06/borrador-v0.1 y QA PDF nuevo.
Requiere dependencias documentales ya instaladas. Falla si algún destino existe.
"""
from html import escape
import hashlib
import json
from pathlib import Path
import re
import sys

sys.dont_write_bytecode = True
from artefactos_hito import bloques, sin_marcas

ROOT = Path(__file__).resolve().parents[2]
FUENTE = ROOT / "documentos/proyecto/Semana6_AI_Native_SDLC_v0.1.md"
SALIDA = ROOT / "entregas/semana-06/borrador-v0.1"
FIGURAS = ROOT / "arquitectura/semana-06"
QA = ROOT / "evidencias/semana-06/qa_pdf"
NOMBRE = "TISG_E01_S06_AI_Native_SDLC_v0.1_BORRADOR"


def limpio(texto):
    return sin_marcas(re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", texto))


def main():
    import pymupdf
    from PIL import Image, ImageDraw
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image as RLImage

    for destino in [SALIDA, FIGURAS, QA]:
        if destino.exists():
            raise FileExistsError(f"Destino existente, no sobrescribir: {destino}")
    for destino in [SALIDA, FIGURAS, QA]:
        destino.mkdir(parents=True, exist_ok=False)
    cajas = [
        ("01 REQUISITOS", "Hito I, supuestos y contrato", "A1: proponer; equipo confirma"),
        ("02 DISENO", "ADR y arquitectura v1.2", "A1: alternativas; equipo decide"),
        ("03 IMPLEMENTACION", "POC-02: Python, solo lectura", "A2: generar; revision pendiente"),
        ("04 PRUEBAS", "36 casos + 12 controles", "A2: ejecutar; no independiente"),
        ("05 REVISION / SEGURIDAD", "Hashes, limites, datos sinteticos", "Persona acepta o rechaza"),
        ("06 INTEGRACION", "Git + verificador solo lectura", "A2: preparar; mantenedor autoriza"),
        ("07 DESPLIEGUE", "No realizado ni autorizado", "A0: control humano; no ERP"),
        ("08 MANTENER / DOCUMENTAR", "Logs, contrato y borradores", "A1/A2: revisar cada cambio"),
    ]
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="760" viewBox="0 0 1200 760">',
           '<rect width="1200" height="760" fill="#f4f7fc"/>',
           '<g font-family="Arial, sans-serif"><text x="45" y="48" font-size="28" fill="#153252">SDLC academico - Grupo 1 / Semana 6</text>',
           '<text x="45" y="80" font-size="18" fill="#34455b">Hechos del repositorio; SDLC empresarial NO DETERMINADO en todas las fases</text>']
    for i, (titulo, hecho, control) in enumerate(cajas):
        x = 45 + (i % 2) * 575
        y = 115 + (i // 2) * 132
        color = "#e5f1ed" if i in [2, 3] else "#ffffff"
        svg += [f'<rect x="{x}" y="{y}" width="535" height="112" rx="12" fill="{color}" stroke="#52718c"/>',
                f'<text x="{x+18}" y="{y+30}" font-size="19" font-weight="bold" fill="#153252">{escape(titulo)}</text>',
                f'<text x="{x+18}" y="{y+59}" font-size="17" fill="#34455b">{escape(hecho)}</text>',
                f'<text x="{x+18}" y="{y+87}" font-size="16" fill="#34455b">{escape(control)}</text>']
    svg += ['<text x="45" y="674" font-size="18" fill="#153252">Lectura: 01 a 08; iterar ante hallazgos. Seguridad y documentacion son transversales.</text>',
            '<text x="45" y="707" font-size="18" fill="#a13b28">Puerta humana: no publicar ni desplegar hasta revision y autorizacion reales.</text>',
            '<text x="45" y="738" font-size="15" fill="#34455b">Verde: codigo y pruebas ejecutados en S6. Resto: antecedentes o propuestas. BORRADOR / IA.</text></g></svg>']
    (FIGURAS / "mapa-sdlc.svg").write_text("\n".join(svg), encoding="utf-8")
    with pymupdf.open(FIGURAS / "mapa-sdlc.svg") as figura:
        (FIGURAS / "mapa-sdlc.pdf").write_bytes(figura.convert_to_pdf())
    with pymupdf.open(FIGURAS / "mapa-sdlc.pdf") as figura:
        figura[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(FIGURAS / "mapa-sdlc.png")

    texto = FUENTE.read_text(encoding="utf-8")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.6)
    estilo = doc.styles["Normal"]
    estilo.font.name, estilo.font.size = "Arial", Pt(10)
    estilo.paragraph_format.space_after = Pt(5)
    sec.footer.paragraphs[0].text = "Grupo 1 | S6 v0.1 BORRADOR | Pendiente de validación humana"
    doc.core_properties.title = "Semana 6 AI-Native SDLC — BORRADOR v0.1"
    doc.core_properties.author = "Grupo 1 — generado con IA; revisión pendiente"

    fontdir = Path("C:/Windows/Fonts")
    pdfmetrics.registerFont(TTFont("S6Arial", str(fontdir / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("S6ArialB", str(fontdir / "arialbd.ttf")))
    estilos = getSampleStyleSheet()
    estilos.add(ParagraphStyle(name="S6", fontName="S6Arial", fontSize=9, leading=13, spaceAfter=7, splitLongWords=True))
    estilos.add(ParagraphStyle(name="S6Cell", parent=estilos["S6"], fontSize=7.3, leading=10, spaceAfter=0))
    estilos.add(ParagraphStyle(name="S6H", fontName="S6ArialB", fontSize=12, leading=16, spaceBefore=12, spaceAfter=7, keepWithNext=True))
    historia = []
    for tipo, valor in bloques(texto):
        if tipo == "titulo":
            nivel, contenido = valor
            doc.add_heading(limpio(contenido), level=min(nivel, 3))
            historia.append(Paragraph(escape(limpio(contenido)), estilos["S6H"]))
        elif tipo == "tabla":
            tabla = doc.add_table(rows=1, cols=len(valor[0]))
            tabla.style = "Light Shading Accent 1"
            for celda, contenido in zip(tabla.rows[0].cells, valor[0]):
                celda.text = limpio(contenido)
            tabla.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for fila in valor[1:]:
                for celda, contenido in zip(tabla.add_row().cells, fila):
                    celda.text = limpio(contenido)
            for fila in tabla.rows:
                fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for celda in fila.cells:
                    for parrafo in celda.paragraphs:
                        for run in parrafo.runs:
                            run.font.size = Pt(8)
            celdas = [[Paragraph(escape(limpio(c)), estilos["S6Cell"]) for c in fila] for fila in valor]
            t = LongTable(celdas, colWidths=[493 / len(valor[0])] * len(valor[0]), repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce8f4")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#9bafc4")), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            historia.extend([t, Spacer(1, 8)])
        elif tipo == "imagen":
            alt, ruta = valor
            imagen = (FUENTE.parent / ruta).resolve()
            doc.add_picture(str(imagen), width=Cm(17.4))
            doc.add_paragraph(alt, style="Caption")
            historia.append(RLImage(str(imagen), width=493, height=493 * 760 / 1200))
            historia.append(Paragraph(escape(alt), estilos["S6Cell"]))
        else:
            doc.add_paragraph(limpio(valor))
            historia.append(Paragraph(escape(limpio(valor)), estilos["S6"]))
    doc.save(SALIDA / (NOMBRE + ".docx"))

    def pie(canvas, documento):
        canvas.setFont("S6Arial", 8)
        canvas.drawString(45, 24, "Grupo 1 | S6 v0.1 BORRADOR | Pendiente de validación humana")
        canvas.drawRightString(A4[0] - 45, 24, str(documento.page))

    pdf = SALIDA / (NOMBRE + ".pdf")
    SimpleDocTemplate(str(pdf), pagesize=A4, leftMargin=51, rightMargin=51, topMargin=43, bottomMargin=43,
                      title="Semana 6 AI-Native SDLC — BORRADOR", author="Grupo 1 / revisión humana pendiente").build(historia, onFirstPage=pie, onLaterPages=pie)
    paginas = []
    with pymupdf.open(pdf) as documento:
        for i, pagina in enumerate(documento):
            ruta = QA / f"pagina-{i+1:02}.png"
            pagina.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(ruta)
            paginas.append({"pagina": i+1, "caracteres": len(pagina.get_text()), "fuera_pagina": sum(1 for b in pagina.get_text("blocks") if b[0] < 0 or b[1] < 0 or b[2] > pagina.rect.width + 1 or b[3] > pagina.rect.height + 1)})
    miniaturas = []
    for p in sorted(QA.glob("pagina-*.png")):
        im = Image.open(p).convert("RGB")
        im.thumbnail((300, 425))
        miniaturas.append(im.copy())
    contacto = Image.new("RGB", (3 * 320, ((len(miniaturas)+2)//3)*455), "#d8dfe8")
    dibujar = ImageDraw.Draw(contacto)
    for i, im in enumerate(miniaturas):
        x, y = (i % 3)*320+10, (i//3)*455+22
        contacto.paste(im, (x, y))
        dibujar.text((x, y-18), f"Pagina {i+1}", fill="black")
    contacto.save(QA / "contacto.png")
    (QA / "estructura.json").write_text(json.dumps({"paginas": paginas, "docx_render_word": False, "revision_visual_IA": "PENDIENTE DE INSPECCION"}, indent=2), encoding="utf-8")
    manifiesto = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in [FUENTE, *SALIDA.iterdir(), *FIGURAS.iterdir()] if p.is_file()}
    (SALIDA / "manifiesto.json").write_text(json.dumps(manifiesto, indent=2), encoding="utf-8")
    print(json.dumps({"pdf": pdf.as_posix(), "paginas": len(paginas), "docx": True, "borrador": True}))


if __name__ == "__main__":
    main()
