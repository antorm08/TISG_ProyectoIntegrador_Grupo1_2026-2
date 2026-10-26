#!/usr/bin/env python3
"""Comprobación técnica v1.2, independiente del generador de formatos.

Solo lectura por defecto. --guardar escribe exclusivamente evidencia v1.2.
No ejecuta PoC, no regenera entregables ni acredita revisión humana.
"""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
import pymupdf

ROOT = Path(__file__).resolve().parents[2]
FUENTE = ROOT / "documentos/proyecto/Semana5_Hito_I_v1.2.md"
ENTREGA = ROOT / "entregas/semana-05/revision-v1.2/TISG_E01_S05_HitoI_v1.2_2026-10-25"
EVIDENCIA = ROOT / "evidencias/revision-v1.2"


def normalizar(texto):
    return re.sub(r"\s+", "", str(texto).replace("**", "").replace("`", ""))


def tablas(texto):
    resultado, actual = [], []
    for linea in [*texto.splitlines(), ""]:
        if linea.startswith("|"):
            fila = [c.strip() for c in linea.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in fila):
                actual.append(fila)
        elif actual:
            assert all(len(f) == len(actual[0]) for f in actual)
            resultado.append(actual)
            actual = []
    return resultado


def verificar():
    resultados = {}
    antes = json.loads((EVIDENCIA / "integridad_antes.json").read_text(encoding="utf-8"))
    diferencias = [p for p, h in antes["archivos"].items()
                   if not (ROOT / p).is_file() or hashlib.sha256((ROOT / p).read_bytes()).hexdigest() != h]
    assert set(diferencias) <= {"README.md"}, diferencias
    extras = json.loads((EVIDENCIA / "fuentes_locales_antes.json").read_text(encoding="utf-8"))
    assert all(hashlib.sha256((ROOT / p).read_bytes()).hexdigest() == h for p, h in extras.items())
    assert subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT).decode().strip() == antes["head"]
    resultados["historicos_y_auditoria_original_intactos"] = True
    resultados["archivos_previos_comparados"] = len(antes["archivos"]) + len(extras)
    resultados["diferencias_previas_permitidas"] = diferencias
    resultados["head_sin_cambios"] = antes["head"]

    texto = FUENTE.read_text(encoding="utf-8")
    ts = tablas(texto)
    doc = Document(str(ENTREGA) + ".docx")
    assert len(doc.tables) == len(ts)
    for t, esperado in zip(doc.tables, ts):
        real = [[c.text for c in f.cells] for f in t.rows]
        assert [[normalizar(c) for c in f] for f in real] == [[normalizar(c) for c in f] for f in esperado]
    imagenes = re.findall(r"!\[(.*?)\]\((.*?)\)", texto)
    assert len(imagenes) == len(doc.inline_shapes) == 4
    for _, ruta in imagenes:
        assert (FUENTE.parent / ruta).resolve().is_file()
    for fuente, destino in [
        (FUENTE, Path(str(ENTREGA) + ".docx")),
        (ROOT / "documentos/proyecto/Semana2_MaterialComplementario_v1.2.md", ROOT / "entregas/semana-02/revision-v1.2/MaterialComplementario_S2_v1.2.docx"),
    ]:
        d = Document(destino)
        contenido = normalizar("\n".join(p.text for p in d.paragraphs))
        for ln in fuente.read_text(encoding="utf-8").splitlines():
            if ln and not ln.startswith(("|", "![", "---")):
                esperado = re.sub(r"^#{1,6}\s+", "", ln)
                assert normalizar(esperado) in contenido, (destino.name, ln[:100])
    resultados["tablas_docx_identicas"] = len(ts)
    resultados["imagenes_docx"] = 4
    resultados["contenido_parrafos_docx"] = "OK; no renderizado Word"

    with pymupdf.open(str(ENTREGA) + ".pdf") as pdf:
        contenido = normalizar("\n".join(p.get_text(clip=pymupdf.Rect(0, 0, p.rect.width, p.rect.height - 40)) for p in pdf))
        faltantes = [c for t in ts for f in t for c in f if normalizar(c) not in contenido]
        assert not faltantes, faltantes[:5]
        for ln in texto.splitlines():
            if ln and not ln.startswith(("|", "![", "---")):
                esperado = re.sub(r"^#{1,6}\s+", "", ln)
                assert normalizar(esperado) in contenido, ln[:120]
        assert sum(len(p.get_images()) for p in pdf) == 4
        fuera = []
        for i, p in enumerate(pdf):
            for b in p.get_text("dict")["blocks"]:
                for linea in b.get("lines", []):
                    for span in linea["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        if x0 < -1 or y0 < -1 or x1 > p.rect.width + 1 or y1 > p.rect.height + 1:
                            fuera.append([i + 1, span["text"]])
        assert not fuera, fuera[:5]
        resultados["pdf_paginas"] = len(pdf)
        resultados["pdf_tablas_parrafos_figuras"] = "Contenido íntegro; revisión visual por separado"
        resultados["texto_fuera_de_pagina"] = fuera

    for p in [Path(str(ENTREGA) + ".docx"), ROOT / "entregas/semana-02/revision-v1.2/MaterialComplementario_S2_v1.2.docx", ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.2.xlsx"]:
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None
            for nombre in z.namelist():
                if nombre.endswith(".xml"):
                    ET.fromstring(z.read(nombre))
    resultados["integridad_zip_xml"] = "OK"

    w = load_workbook(ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.2.xlsx", data_only=False, keep_links=True)
    old = load_workbook(ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.1.xlsx", data_only=False, keep_links=True)
    for nombre in ("Contribuciones", "Fuentes_Historicas", "Supuestos_Economicos"):
        assert list(w[nombre].values) == list(old[nombre].values), nombre
    for nombre, origen in [("Madurez_Hist_v11", "Madurez_SocioTecnica"), ("Fuentes_Hist_v11", "Fuentes"), ("Hipotesis_Hist_v11", "Hipotesis"), ("Supuestos_Hist_v11", "Supuestos")]:
        assert list(w[nombre].values) == list(old[origen].values), nombre
    formulas_antes = {(s.title, c.coordinate): c.value for s in old for f in s for c in f if c.data_type == "f"}
    assert all(w[n][celda].value == formula for (n, celda), formula in formulas_antes.items())
    assert [r[0] for r in list(w["Decisiones"].values)[1:]] == [f"D-{i:02}" for i in range(1, 23)]
    assert [r[0] for r in list(w["Supuestos"].values)[1:]] == [f"S-{i:02}" for i in range(1, 27)]
    for nombre in ["ADR", "Radar_Data", "Canvas_Data", "Propuestas", "Stakeholders", "Madurez_SocioTecnica", "Skills Gap", "Impact_Canvas", "Hipotesis", "Deuda_Trazabilidad", "Dependencias", "Instrumento_Madurez", "Metricas", "Umbrales", "Datos_Relaciones", "Topologia", "Revision_Humana", "Coherencia"]:
        filas = [[normalizar(v) for v in f] for f in w[nombre].values]
        assert any(filas == [[normalizar(c) for c in f] for f in t] for t in ts), nombre
    original_fuentes = tablas((ROOT / "fuentes/revision-v1.2/VERIFICACION_FUENTES.md").read_text(encoding="utf-8"))
    for nombre in ("Fuentes", "CAPECE_Correcciones"):
        filas = [[normalizar(v) for v in f] for f in w[nombre].values]
        assert any(filas == [[normalizar(c) for c in f] for f in t] for t in original_fuentes), nombre
    resultados["hojas_excel"] = len(w.sheetnames)
    resultados["formulas_originales_preservadas"] = len(formulas_antes)
    resultados["excel_sin_recalculo_simulado"] = True
    resultados["contribuciones_originales_conservadas"] = w["Contribuciones"].max_row - 1
    old.close()
    w.close()

    figuras = sorted((ROOT / "arquitectura/revision-v1.2").glob("*.svg"))
    laminas = sorted((ROOT / "entregas/semana-02/revision-v1.2").glob("lamina-*.svg"))
    assert len(figuras) == 4 and len(laminas) == 5
    for p in figuras + laminas:
        ET.fromstring(p.read_bytes())
        with pymupdf.open(p.with_suffix(".pdf")) as pdf:
            assert len(pdf) == 1
            assert "v1.2" in pdf[0].get_text()
    with pymupdf.open(ROOT / "entregas/semana-02/revision-v1.2/Presentacion_S2_v1.2.pdf") as p:
        assert len(p) == 5
        t = p[0].get_text()
        assert "compradores" in t
        assert "Contratantes" in t and "Canales" in t and "Segmentos" in t
    resultados["figuras_laminas"] = [len(figuras), len(laminas)]
    resultados["experimento_v12"] = "NO EJECUTADO: solo diseño"
    resultados["revision_humana"] = "PENDIENTE"

    post = ROOT / "AUDITORIA_POST_CORRECCIONES_S01_S05_v1.2.md"
    if post.exists():
        checklist = post.read_text(encoding="utf-8").split("## 4. Checklist completo de 69 criterios", 1)[1].split("## 5. Justificación", 1)[0]
        filas = [ln.split("|") for ln in checklist.splitlines() if re.match(r"^\| [1-5]\.\d{2} \|", ln)]
        assert len(filas) == 69
        originales = [ln.split("|") for ln in (ROOT / "AUDITORIA_CUMPLIMIENTO_S01_S05.md").read_text(encoding="utf-8").splitlines() if re.match(r"^\| [1-5]\.\d{2} \|", ln)]
        assert [(f[1].strip(), f[2].strip()) for f in filas] == [(f[1].strip(), f[2].strip()) for f in originales]
        estados = Counter(f[3].strip() for f in filas)
        assert estados == {"CUMPLE": 67, "PARCIAL": 2}, estados
        resultados["auditoria_69_criterios"] = dict(estados)
    return resultados


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--guardar", action="store_true")
    args = parser.parse_args()
    r = verificar()
    salida = json.dumps(r, ensure_ascii=False, indent=2)
    if args.guardar:
        (EVIDENCIA / "verificacion_tecnica.json").write_text(salida + "\n", encoding="utf-8")
    print(salida)
