"""Verificación S6 de SOLO LECTURA. No ejecuta PoC, generadores ni red."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import unquote, urlsplit
import zipfile

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
POC = ROOT / "poc/POC-02-ai-native-sdlc"


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def cargar(p):
    return json.loads(p.read_text(encoding="utf-8"))


def verificar():
    import pymupdf
    from docx import Document

    comprobaciones = []

    def check(condicion, nombre):
        if not condicion:
            raise ValueError(nombre)
        comprobaciones.append(nombre)

    inicial = cargar(ROOT / "evidencias/semana-06/integridad_antes.json")
    iguales = [n for n, h in inicial["archivos"].items() if (ROOT / n).is_file() and sha(ROOT / n) == h]
    cambiados = sorted(set(inicial["archivos"]) - set(iguales))
    check(cambiados == ["README.md"], "Históricos intactos; única edición previa: README.md")
    base_readme = subprocess.check_output(["git", "show", inicial["head_base"] + ":README.md"], cwd=ROOT).decode("utf-8").replace("\r\n", "\n")
    check((ROOT / "README.md").read_text(encoding="utf-8").startswith(base_readme.rstrip() + "\n"), "README conserva todo el contenido previo y añade navegación")
    check(subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT).decode().strip() == inicial["head_base"], "Sin commits nuevos")
    congelados = cargar(ROOT / "evidencias/semana-06/congelacion_previa.json")
    for n, h in congelados["sha256"].items():
        check(sha(ROOT / n) == h, "Congelación: " + n)
    primera = ROOT / "evidencias/semana-06/primera_version"
    for n, h in cargar(primera / "manifiesto.json")["sha256"].items():
        check(sha(primera / Path(n).name) == h == sha(ROOT / n), "Primera versión sin correcciones: " + n)
    corrida = POC / "resultados/S6-001"
    for n, h in cargar(corrida / "manifiesto.json").items():
        check(sha(corrida / n) == h, "Corrida: " + n)
    r = cargar(corrida / "resultados.json")
    casos = r["casos"]
    esperados = cargar(POC / "pruebas/casos.json")["casos"]
    check(len(casos) == len(esperados) == 36, "36 casos distintos ejecutados")
    check({c["id"] for c in casos} == {c["id"] for c in esperados}, "Identidades del corpus completas")
    por_id = {c["id"]: c for c in esperados}
    for c in casos:
        e = por_id[c["id"]]
        check(c["valido_esperado"] == e["valido"] and Counter(c["codigos_esperados"]) == Counter(e["codigos"]), "Esperados congelados: " + c["id"])
        check(c["correcto"] and c["entradas_intactas"] and c["valido_observado"] == e["valido"] and Counter(c["codigos_observados"]) == Counter(e["codigos"]), "Caso correcto e intacto: " + c["id"])
    m = r["metricas"]
    validos = [c for c in casos if c["valido_esperado"]]
    invalidos = [c for c in casos if not c["valido_esperado"]]
    check(len(validos) == m["validos"] == 8 and len(invalidos) == m["invalidos"] == 28, "Denominadores Q2/Q3: 28/8")
    check(m["Q1_porcentaje"] == 100 * sum(c["correcto"] for c in casos) / len(casos), "Q1 recomputado")
    check(m["Q2_falsas_aceptaciones"] == sum(c["valido_observado"] for c in invalidos) == 0, "Q2 recomputado")
    check(m["Q3_falsos_rechazos"] == sum(not c["valido_observado"] for c in validos) == 0, "Q3 recomputado")
    check(len(r["tests"]) == m["tests_total"] == 48 and all(t["estado"] == "OK" for t in r["tests"]) and m["tests_fallos"] == m["tests_errores"] == 0, "48 tests OK, sin fallos/errores")
    check("Ran 48 tests" in (corrida / "unittest.log").read_text(encoding="utf-8") and (corrida / "unittest.log").read_text(encoding="utf-8").rstrip().endswith("OK"), "Log unittest coherente")
    check(r["aceptacion_tecnica"] and r["originales_intactos"] and r["cli_exit"] == 0 and r["revision_humana"] == "PENDIENTE", "Aceptación técnica separada de revisión humana")
    check(all(m[n] == "NO DETERMINADO" for n in ["M1_desarrollo", "M2_revision_humana", "M4_ahorro"]), "Sin métricas humanas ficticias")
    check(cargar(corrida / "cli_stdout.json") == {"version_contrato": "1.0", "estado": "VALIDO", "archivos": {"costos": {"filas": 35}, "metrados": {"filas": 35}}, "errores": []}, "CLI válida sobre originales 35/35")
    entregas = ROOT / "entregas/semana-06/borrador-v0.1"
    for n, h in cargar(entregas / "manifiesto.json").items():
        check(sha(ROOT / n) == h, "Exportación: " + n)
    pdf = next(entregas.glob("*.pdf"))
    with pymupdf.open(pdf) as d:
        check(len(d) == 8 and all("BORRADOR" in p.get_text() for p in d), "PDF 8 páginas y estado BORRADOR en todas")
        check(all(p.get_text().strip() for p in d), "PDF sin páginas vacías")
    docx = next(entregas.glob("*.docx"))
    with zipfile.ZipFile(docx) as z:
        check(z.testzip() is None, "DOCX ZIP íntegro")
    doc = Document(docx)
    check(len(doc.tables) == 6 and "BORRADOR" in "\n".join(p.text for p in doc.paragraphs), "DOCX seis tablas y estado borrador")
    reporte = (ROOT / "documentos/proyecto/Semana6_AI_Native_SDLC_v0.1.md").read_text(encoding="utf-8")
    check(len(re.findall(r"^## \d+\.", reporte, re.M)) == 25, "Informe con 25 secciones")
    for token in ["1,448975", "0,178095", "36/36", "48/48", "0/28", "0/8", "NO DETERMINADO", "Souppaya", "Becker", "DORA", "Python Software Foundation"]:
        check(token in reporte, "Informe/evidencias: " + token)
    no_versionados = subprocess.check_output(["git", "ls-files", "--others", "--exclude-standard", "-z"], cwd=ROOT).decode().split("\0")
    nuevos = [ROOT / n for n in no_versionados if n and (ROOT / n).is_file()]
    enlaces = 0
    for p in nuevos + [ROOT / "README.md"]:
        if p.suffix.lower() == ".md":
            text = p.read_text(encoding="utf-8")
            for destino in re.findall(r"!?\[[^\]]*\]\(([^)]+)\)", text):
                url = urlsplit(destino)
                if url.scheme or not url.path:
                    continue
                check((p.parent / unquote(url.path)).resolve().exists(), "Enlace: " + p.relative_to(ROOT).as_posix() + " -> " + destino)
                enlaces += 1
    patron = re.compile(r"-----BEGIN (?:RSA |OPENSSH |EC )?PRIVATE KEY-----|\bgh[pousr]_[A-Za-z0-9]{30,}|\bAKIA[A-Z0-9]{16}\b")
    for p in nuevos:
        if p.suffix.lower() in {".md", ".py", ".json", ".csv", ".log", ".svg"}:
            check(not patron.search(p.read_text(encoding="utf-8")), "Sin patrones de credenciales: " + p.relative_to(ROOT).as_posix())
    return {"estado": "OK", "comprobaciones": len(comprobaciones), "previos_iguales": len(iguales), "previos_modificados": cambiados,
            "enlaces_locales_verificados": enlaces, "casos": len(casos), "tests": len(r["tests"]), "revision_humana": "PENDIENTE", "limite_escaneo": "Patrones acotados, no auditoría de seguridad", "detalle": comprobaciones}


if __name__ == "__main__":
    try:
        print(json.dumps(verificar(), ensure_ascii=False, indent=2))
    except Exception as exc:
        print(json.dumps({"estado": "ERROR", "detalle": str(exc)}, ensure_ascii=False))
        raise SystemExit(1)
