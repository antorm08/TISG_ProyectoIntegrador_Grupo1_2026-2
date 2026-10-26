"""Exportación editorial v0.2 desde Markdown. Destino de construcción NUEVO por ID.

No publica ni sobrescribe v0.1 o PDF/DOCX finales. Renderiza todas las páginas.
"""
from html import escape
import argparse
import hashlib
import json
from pathlib import Path
import re
import sys

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "documentos/proyecto/Semana6_AI_Native_SDLC_v0.2.md"
NAME = "TISG_E01_S06_AI_Native_SDLC_v0.2_BORRADOR"


def tokens(texto):
    lines = texto.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].strip()
        i += 1
        if not s:
            continue
        if s == "<!-- pagebreak -->":
            yield "break", None
        elif s.startswith("```"):
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            i += 1
            yield "code", "\n".join(block)
        elif s.startswith("<!--"):
            continue
        elif s.startswith("|"):
            table = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                table.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in row.strip("|").split("|")] for row in table]
            yield "table", [row for row in rows if not all(re.fullmatch(r":?-+:?", c) for c in row)]
        elif s.startswith("#"):
            yield "heading", (len(s)-len(s.lstrip("#")), s.lstrip("# "))
        elif re.fullmatch(r"!\[(.*?)\]\((.*?)\)", s):
            yield "image", re.fullmatch(r"!\[(.*?)\]\((.*?)\)", s).groups()
        elif s.startswith("> "):
            yield "callout", s[2:]
        else:
            yield "text", s


def plain(s):
    s = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", s)
    return s.replace("**", "").replace("`", "").replace("*", "")


def rich(s):
    s = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", s)
    s = escape(s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"`([^`]+)`", r'<font name="Mono">\1</font>', s)
    s = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", s)
    return s


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--id", required=True)
    args = parser.parse_args()
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", args.id):
        parser.error("ID inválido")
    out = ROOT / "evidencias/semana-06/revision-v0.2" / args.id
    out.mkdir(parents=True, exist_ok=False)
    import pymupdf
    from PIL import Image, ImageDraw
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Preformatted, KeepTogether
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from figuras_semana6_v02 import figuras

    for name, file in [("Body", "arial.ttf"), ("BodyB", "arialbd.ttf"), ("BodyI", "ariali.ttf"), ("Mono", "consola.ttf")]:
        pdfmetrics.registerFont(TTFont(name, "C:/Windows/Fonts/" + file))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="BodyB", italic="BodyI", boldItalic="BodyB")
    navy, teal, gray, ink = map(colors.HexColor, ["#102C46", "#087F83", "#EDF3F6", "#263D4D"])
    styles = {
        "text": ParagraphStyle("body", fontName="Body", fontSize=10.2, leading=14.7, textColor=ink, spaceAfter=9),
        "h1": ParagraphStyle("title", fontName="BodyB", fontSize=32, leading=37, textColor=navy, spaceAfter=22),
        "h2": ParagraphStyle("section", fontName="BodyB", fontSize=20, leading=25, textColor=navy, spaceBefore=2, spaceAfter=15, keepWithNext=True),
        "h3": ParagraphStyle("subsection", fontName="BodyB", fontSize=12, leading=16, textColor=teal, spaceBefore=10, spaceAfter=8, keepWithNext=True),
        "cell": ParagraphStyle("cell", fontName="Body", fontSize=9, leading=12, textColor=ink, spaceAfter=0),
        "th": ParagraphStyle("th", fontName="BodyB", fontSize=9, leading=12, textColor=colors.white),
        "caption": ParagraphStyle("caption", fontName="Body", fontSize=8.6, leading=12, textColor=teal, spaceBefore=6, spaceAfter=12),
        "callout": ParagraphStyle("callout", fontName="BodyB", fontSize=10, leading=14, textColor=navy, spaceAfter=0),
        "code": ParagraphStyle("code", fontName="Mono", fontSize=8, leading=10.3, textColor=ink, leftIndent=10, rightIndent=8, spaceBefore=6, spaceAfter=10, backColor=gray, borderPadding=9),
    }
    source = SOURCE.read_text(encoding="utf-8")
    # Snapshot literal, no documento navegable desde el directorio de build.
    (out / (SOURCE.name + ".txt")).write_text(source, encoding="utf-8")
    all_tokens = list(tokens(source))
    figures = figuras()
    story = [Spacer(1, 50)]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for n in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[n].font.name = "Arial"
        doc.styles[n].font.color.rgb = RGBColor.from_string("102C46" if n != "Heading 3" else "087F83")
    sec.footer.paragraphs[0].text = "Grupo 1 | S6 v0.2 BORRADOR | Pendiente de revisión del equipo"
    sec.header.paragraphs[0].text = "UNMSM - FISI   /   Tendencias en Ingeniería de Software y Gestión"
    doc.core_properties.title = "Semana 6 - AI-Native SDLC y productividad - v0.2 BORRADOR"
    doc.core_properties.author = "Grupo 1 / elaboración IA; revisión humana pendiente"
    for kind, value in all_tokens:
        if kind == "break":
            story.append(PageBreak())
            doc.add_page_break()
        elif kind == "heading":
            level, text = value
            story.append(Paragraph(rich(text), styles["h1" if level == 1 else "h2" if level == 2 else "h3"]))
            doc.add_heading(plain(text), level=min(level, 3))
        elif kind == "text":
            story.append(Paragraph(rich(value), styles["text"]))
            p = doc.add_paragraph()
            for part in re.split(r"(\*\*.*?\*\*)", value):
                run = p.add_run(plain(part))
                run.bold = part.startswith("**")
        elif kind == "callout":
            t = Table([[Paragraph(rich(value), styles["callout"])]], colWidths=[490])
            t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), gray), ("LINEBEFORE", (0,0), (0,-1), 3, teal), ("LEFTPADDING", (0,0), (-1,-1), 13), ("RIGHTPADDING", (0,0), (-1,-1), 12), ("TOPPADDING", (0,0), (-1,-1), 10), ("BOTTOMPADDING", (0,0), (-1,-1), 10)]))
            story.extend([t, Spacer(1, 11)])
            doc.add_paragraph(plain(value), style="Quote")
        elif kind == "code":
            codebox = Table([[Preformatted(value, styles["code"])]], colWidths=[490])
            codebox.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), gray), ("LEFTPADDING", (0,0), (-1,-1), 8), ("RIGHTPADDING", (0,0), (-1,-1), 8), ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
            story.extend([codebox, Spacer(1, 9)])
            p = doc.add_paragraph()
            run = p.add_run(value)
            run.font.name, run.font.size = "Consolas", Pt(8)
        elif kind == "table":
            rows = value
            assert all(len(row) == len(rows[0]) for row in rows), rows
            count = len(rows[0])
            widths = {2: [155,335], 3: [115,185,190], 4: [95,145,115,135], 5: [110,70,100,100,110]}.get(count, [490/count]*count)
            if rows[0][0] == "Grupo":
                widths = [145,45,100,100,100]
            t = Table([[Paragraph(rich(c), styles["th" if i == 0 else "cell"]) for c in row] for i, row in enumerate(rows)], colWidths=widths, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), navy), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, gray]), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 9), ("RIGHTPADDING", (0,0), (-1,-1), 9), ("TOPPADDING", (0,0), (-1,-1), 8), ("BOTTOMPADDING", (0,0), (-1,-1), 8), ("LINEBELOW", (0,0), (-1,0), .8, teal)]))
            story.extend([t, Spacer(1, 12)])
            dt = doc.add_table(rows=1, cols=count)
            dt.style = "Light Shading Accent 1"
            for cell, val in zip(dt.rows[0].cells, rows[0]):
                cell.text = plain(val)
            dt.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for row in rows[1:]:
                for cell, val in zip(dt.add_row().cells, row):
                    cell.text = plain(val)
            for row in dt.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9)
        elif kind == "image":
            alt, ruta = value
            imagen = (SOURCE.parent / ruta).resolve()
            key = imagen.stem
            if key in figures:
                story.append(figures[key])
            else:
                from reportlab.platypus import Image as RLImage
                with Image.open(imagen) as im:
                    w, h = im.size
                story.append(RLImage(str(imagen), width=490, height=490*h/w))
            story.append(Paragraph(escape(alt), styles["caption"]))
            doc.add_picture(str(imagen), width=Cm(17.2))
            doc.add_paragraph(alt, style="Caption")

    def margins(canvas, document):
        canvas.saveState()
        width, height = A4
        if document.page == 1:
            canvas.setFillColor(navy)
            canvas.rect(0, 0, 13, height, fill=1, stroke=0)
            canvas.setFillColor(gray)
            canvas.setFont("BodyB", 100)
            canvas.drawRightString(width-44, height-98, "06")
        else:
            canvas.setFont("BodyB", 8)
            canvas.setFillColor(teal)
            canvas.drawString(52, height-29, "UNMSM / FISI / GRUPO 1")
            canvas.setFont("Body", 8)
            canvas.drawRightString(width-52, height-29, "AI-NATIVE SDLC   /   v0.2")
            canvas.setStrokeColor(colors.HexColor("#D9E3E8"))
            canvas.line(52, height-38, width-52, height-38)
        canvas.setFont("Body", 7.5)
        canvas.setFillColor(ink)
        canvas.drawString(52, 25, "BORRADOR - pendiente de revisión del equipo | 26/10/2026")
        canvas.setFont("BodyB", 9)
        canvas.drawRightString(width-52, 25, f"{document.page:02}")
        canvas.restoreState()

    pdf = out / (NAME + ".pdf")
    SimpleDocTemplate(str(pdf), pagesize=A4, leftMargin=52, rightMargin=52, topMargin=57, bottomMargin=48, title=doc.core_properties.title, author=doc.core_properties.author).build(story, onFirstPage=margins, onLaterPages=margins)
    doc.save(out / (NAME + ".docx"))
    render = out / "paginas"
    render.mkdir()
    pages = []
    with pymupdf.open(pdf) as d:
        for i, page in enumerate(d):
            page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(render / f"pagina-{i+1:02}.png")
            blocks = page.get_text("blocks")
            pages.append({"pagina": i+1, "caracteres": len(page.get_text()), "fuera_lienzo": sum(1 for b in blocks if b[0] < 0 or b[1] < 0 or b[2] > page.rect.width+.5 or b[3] > page.rect.height+.5), "texto": page.get_text()})
    contact = Image.new("RGB", (3*320, ((len(pages)+2)//3)*460), "#D8E3E8")
    drawer = ImageDraw.Draw(contact)
    for i in range(len(pages)):
        with Image.open(render / f"pagina-{i+1:02}.png") as im:
            im.thumbnail((300,425))
            x,y=(i%3)*320+10,(i//3)*460+25
            contact.paste(im,(x,y));drawer.text((x,y-18),f"Pagina {i+1}",fill="black")
    contact.save(out / "contacto.png")
    (out / "estructura.json").write_text(json.dumps({"paginas": pages, "docx_renderizado": False, "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest()}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"id": args.id, "paginas": len(pages), "pdf": str(pdf), "fuera_lienzo": sum(p["fuera_lienzo"] for p in pages)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
