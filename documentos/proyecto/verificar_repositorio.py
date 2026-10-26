#!/usr/bin/env python3
"""Verificación ACTUAL, solo lectura y sin conexión de red.

No ejecuta generadores/PoC ni acredita aprobación humana o una fotografía Git.
Puede verificar una exportación sin .git con --root. Reutiliza únicamente los
parsers puros del verificador histórico; nunca llama a su función verificar().
"""

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import sys
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

sys.dont_write_bytecode = True

from docx import Document
from openpyxl import load_workbook
import pymupdf

from verificar_hito_v12 import normalizar, tablas

ROOT = Path(__file__).resolve().parents[2]
EDITABLES_MANTENIMIENTO = {
    ".gitignore", "README.md", "CHANGELOG.md", "CONTRIBUTING.md",
    "arquitectura/as-is/README.md", "arquitectura/to-be/README.md",
    "poc/POC-01-estimacion-historicos/README.md",
}
LOCALES_OPCIONALES = {
    "AUDITORIA_CUMPLIMIENTO_S01_S05.md",
    "documentos/herramientas/AGENT_SKILLS.md",
}


def exigir(condicion, mensaje):
    if not condicion:
        raise ValueError(mensaje)


def sha(ruta):
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def identidad_exportada(ruta, esperado):
    """Distingue bytes iguales de conversión LF/CRLF en un checkout/exportación."""
    if not ruta.is_file():
        return "ausente"
    datos = ruta.read_bytes()
    if hashlib.sha256(datos).hexdigest() == esperado:
        return "bytes"
    if ruta.suffix.lower() in {".md", ".py", ".txt", ".svg", ".json", ".csv"}:
        lf = datos.replace(b"\r\n", b"\n")
        if any(hashlib.sha256(v).hexdigest() == esperado for v in (lf, lf.replace(b"\n", b"\r\n"))):
            return "solo_LF_CRLF"
    return "diferente"


def sin_pie(pdf):
    return normalizar("\n".join(p.get_text(clip=pymupdf.Rect(
        0, 0, p.rect.width, p.rect.height - 40)) for p in pdf))


def verificar_enlaces(root):
    enlaces = []
    for p in sorted(root.rglob("*.md")):
        relativa = p.relative_to(root)
        if any(x in relativa.parts for x in (".git", ".agents", ".opencode")):
            continue
        if relativa.as_posix() in LOCALES_OPCIONALES:
            continue
        texto = re.sub(r"```.*?```", "", p.read_text(encoding="utf-8"), flags=re.S)
        for destino in re.findall(r"!?\[[^\]]*\]\(([^\n]+?)\)", texto):
            url = urlsplit(destino.strip().strip("<>"))
            if url.scheme or not url.path:
                continue
            ruta = (p.parent / unquote(url.path)).resolve()
            exigir(ruta.is_relative_to(root), f"Enlace fuera del repositorio: {relativa}")
            exigir(ruta.exists(), f"Enlace roto en {relativa}: {destino}")
            enlaces.append([relativa.as_posix(), destino])
    return len(enlaces)


def verificar(root):
    root = root.resolve()
    resultados = {"modo": "estado_actual_solo_lectura", "revision_humana": "PENDIENTE"}
    baseline = json.loads((root / "evidencias/mantenimiento-s06/integridad_antes.json").read_text(encoding="utf-8"))
    protegidos = {p: h for p, h in baseline["archivos"].items() if p not in EDITABLES_MANTENIMIENTO}
    identidades = {p: identidad_exportada(root / p, h) for p, h in protegidos.items()}
    diferencias = [p for p, identidad in identidades.items() if identidad in ("ausente", "diferente")]
    exigir(not diferencias, f"Artefactos preservados diferentes del inicio del mantenimiento: {diferencias}")
    resultados["archivos_previos_protegidos"] = len(protegidos)
    resultados["solo_conversion_LF_CRLF"] = [p for p, identidad in identidades.items() if identidad == "solo_LF_CRLF"]
    resultados["base_de_preservacion"] = baseline["head_base"]
    resultados["alcance_preservacion"] = "Bytes locales iniciales, no equivalencia con HEAD ni aprobación histórica"

    guia = root / "proyectoIntegrador.pdf"
    exigir(sha(guia) == baseline["archivos"]["proyectoIntegrador.pdf"], "Guía compartida alterada")
    with pymupdf.open(guia) as d:
        exigir(len(d) == 17, "Guía oficial no tiene 17 páginas")
    local = root / "proyectoIntegrador (1).pdf"
    resultados["guia_canonica"] = "proyectoIntegrador.pdf"
    resultados["copia_local_opcional"] = ("idéntica" if sha(local) == sha(guia) else "DIFIERE: revisar") if local.exists() else "ausente; no requerida"

    fuente = root / "documentos/proyecto/Semana5_Hito_I_v1.2.md"
    texto = fuente.read_text(encoding="utf-8")
    ts = tablas(texto)
    paquete = root / "entregas/semana-05/revision-v1.2/TISG_E01_S05_HitoI_v1.2_2026-10-25"
    doc = Document(str(paquete) + ".docx")
    exigir(len(doc.tables) == len(ts) == 23, "Cantidad de tablas DOCX")
    for t, esperado in zip(doc.tables, ts):
        real = [[normalizar(c.text) for c in f.cells] for f in t.rows]
        exigir(real == [[normalizar(c) for c in f] for f in esperado], "Tabla DOCX distinta del Markdown")
    exigir(len(doc.inline_shapes) == 4, "Figuras DOCX")
    fuente_s2 = root / "documentos/proyecto/Semana2_MaterialComplementario_v1.2.md"
    doc_s2 = root / "entregas/semana-02/revision-v1.2/MaterialComplementario_S2_v1.2.docx"
    for origen, destino in [(fuente, Path(str(paquete) + ".docx")), (fuente_s2, doc_s2)]:
        d = Document(destino)
        contenido = normalizar("\n".join(p.text for p in d.paragraphs))
        for ln in origen.read_text(encoding="utf-8").splitlines():
            if ln and not ln.startswith(("|", "![", "---")):
                exigir(normalizar(re.sub(r"^#{1,6}\s+", "", ln)) in contenido, f"Párrafo DOCX ausente: {origen.name}")
    with pymupdf.open(str(paquete) + ".pdf") as pdf:
        contenido = sin_pie(pdf)
        for celda in (c for t in ts for f in t for c in f):
            exigir(normalizar(celda) in contenido, "Celda ausente del PDF")
        for ln in texto.splitlines():
            if ln and not ln.startswith(("|", "![", "---")):
                exigir(normalizar(re.sub(r"^#{1,6}\s+", "", ln)) in contenido, "Párrafo ausente del PDF")
        exigir(sum(len(p.get_images()) for p in pdf) == 4, "Figuras PDF")
        for p in pdf:
            for b in p.get_text("dict")["blocks"]:
                for linea in b.get("lines", []):
                    for span in linea["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        exigir(x0 >= -1 and y0 >= -1 and x1 <= p.rect.width + 1 and y1 <= p.rect.height + 1, "Texto fuera de página PDF")
        resultados["pdf_paginas"] = len(pdf)

    libro = root / "gestion/registros/Registros_y_Radar_Semana5_v1.2.xlsx"
    for p in [Path(str(paquete) + ".docx"), doc_s2, libro]:
        with zipfile.ZipFile(p) as z:
            exigir(z.testzip() is None, f"ZIP inválido: {p.name}")
            for n in z.namelist():
                if n.endswith(".xml"):
                    ET.fromstring(z.read(n))
    w = load_workbook(libro, data_only=False, keep_links=True)
    old = load_workbook(root / "gestion/registros/Registros_y_Radar_Semana5_v1.1.xlsx", data_only=False, keep_links=True)
    try:
        exigir(len(w.sheetnames) == 34, "Cantidad de hojas XLSX")
        for nombre in ("Contribuciones", "Fuentes_Historicas", "Supuestos_Economicos"):
            exigir(list(w[nombre].values) == list(old[nombre].values), f"Histórico Excel alterado: {nombre}")
        for destino, origen in [("Madurez_Hist_v11", "Madurez_SocioTecnica"), ("Fuentes_Hist_v11", "Fuentes"), ("Hipotesis_Hist_v11", "Hipotesis"), ("Supuestos_Hist_v11", "Supuestos")]:
            exigir(list(w[destino].values) == list(old[origen].values), f"Copia histórica Excel: {destino}")
        formulas = {(s.title, c.coordinate): c.value for s in old for f in s for c in f if c.data_type == "f"}
        exigir(all(w[n][c].value == v for (n, c), v in formulas.items()), "Fórmulas previas alteradas")
        for nombre in ("ADR", "Radar_Data", "Canvas_Data", "Propuestas", "Stakeholders", "Madurez_SocioTecnica", "Skills Gap", "Impact_Canvas", "Hipotesis", "Deuda_Trazabilidad", "Dependencias", "Instrumento_Madurez", "Metricas", "Umbrales", "Datos_Relaciones", "Topologia", "Revision_Humana", "Coherencia"):
            filas = [[normalizar(v) for v in f] for f in w[nombre].values]
            exigir(any(filas == [[normalizar(c) for c in f] for f in t] for t in ts), f"Tabla Excel distinta: {nombre}")
        tf = tablas((root / "fuentes/revision-v1.2/VERIFICACION_FUENTES.md").read_text(encoding="utf-8"))
        for nombre in ("Fuentes", "CAPECE_Correcciones"):
            filas = [[normalizar(v) for v in f] for f in w[nombre].values]
            exigir(any(filas == [[normalizar(c) for c in f] for f in t] for t in tf), f"Fuentes Excel distintas: {nombre}")
        exigir([r[0] for r in list(w["Decisiones"].values)[1:]] == [f"D-{i:02}" for i in range(1, 23)], "IDs decisiones")
        exigir([r[0] for r in list(w["Supuestos"].values)[1:]] == [f"S-{i:02}" for i in range(1, 27)], "IDs supuestos")
        resultados.update(tablas_docx=23, figuras_docx=4, hojas_excel=34, contribuciones_preservadas=w["Contribuciones"].max_row - 1, formulas_previas=len(formulas))
    finally:
        w.close()
        old.close()

    figuras = sorted((root / "arquitectura/revision-v1.2").glob("*.svg"))
    laminas = sorted((root / "entregas/semana-02/revision-v1.2").glob("lamina-*.svg"))
    exigir(len(figuras) == 4 and len(laminas) == 5, "Cantidad de figuras/láminas")
    for p in figuras + laminas:
        ET.fromstring(p.read_bytes())
        exigir(p.with_suffix(".png").is_file(), "Exportación PNG ausente")
        with pymupdf.open(p.with_suffix(".pdf")) as pdf:
            exigir(len(pdf) == 1 and "v1.2" in pdf[0].get_text(), "PDF de figura no corresponde a v1.2")
    with pymupdf.open(root / "entregas/semana-02/revision-v1.2/Presentacion_S2_v1.2.pdf") as pdf:
        exigir(len(pdf) == 5, "Presentación S2 no tiene cinco láminas")
    post = (root / "AUDITORIA_POST_CORRECCIONES_S01_S05_v1.2.md").read_text(encoding="utf-8")
    checklist = post.split("## 4. Checklist completo de 69 criterios", 1)[1].split("## 5. Justificación", 1)[0]
    filas = [ln.split("|") for ln in checklist.splitlines() if re.match(r"^\| [1-5]\.\d{2} \|", ln)]
    ids = [f[1].strip() for f in filas]
    exigir(ids == [f"{sem}.{i:02}" for sem, cantidad in enumerate((13, 11, 11, 13, 21), 1) for i in range(1, cantidad + 1)], "IDs del checklist histórico")
    exigir(Counter(f[3].strip() for f in filas) == {"CUMPLE": 67, "PARCIAL": 2}, "Conteo del informe histórico distinto; no es una reevaluación académica")
    inicial = root / "AUDITORIA_CUMPLIMIENTO_S01_S05.md"
    if inicial.exists():
        originales = [ln.split("|") for ln in inicial.read_text(encoding="utf-8").splitlines() if re.match(r"^\| [1-5]\.\d{2} \|", ln)]
        exigir([(f[1].strip(), f[2].strip()) for f in filas] == [(f[1].strip(), f[2].strip()) for f in originales], "Nombres/IDs distintos de auditoría inicial opcional")
    resultados["auditoria_inicial_opcional"] = "contrastada localmente" if inicial.exists() else "ausente; sin contraste de nombres con el original"
    for relativa in ("arquitectura/as-is/README.md", "arquitectura/to-be/README.md", "poc/POC-01-estimacion-historicos/README.md"):
        exigir("Semana5_Hito_I_v1.2.md" in (root / relativa).read_text(encoding="utf-8"), f"Navegación sin v1.2: {relativa}")
    resultados["enlaces_locales"] = verificar_enlaces(root)
    resultados["limites"] = "Sin ejecución experimental, render Office, recálculo, revisión humana ni nueva auditoría académica"
    return resultados


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT, help="También admite una exportación sin .git")
    args = parser.parse_args()
    try:
        print(json.dumps(verificar(args.root), ensure_ascii=False, indent=2))
    except (ValueError, OSError, KeyError, IndexError, zipfile.BadZipFile, ET.ParseError) as error:
        print(json.dumps({"estado": "ERROR", "detalle": str(error)}, ensure_ascii=False), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
