#!/usr/bin/env python3
"""Verifica consistencia técnica de la revisión; no aprobación académica.

No conecta al ERP ni hace solicitudes HTTP. Guarda la salida real de la PoC
en evidencias/revision-v1.1 sin reescribir su evidencia histórica.
"""

from pathlib import Path
import csv
import json
import os
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from openpyxl import load_workbook
import pymupdf

from artefactos_hito import DESTINO, FUENTE, NOMBRE, ROOT, bloques, leer_tablas, sin_marcas

EVIDENCIAS = ROOT / "evidencias/revision-v1.1"


def normalizar(texto):
    return re.sub(r"\s+", "", sin_marcas(texto))


def verificar():
    EVIDENCIAS.mkdir(parents=True, exist_ok=True)
    texto = FUENTE.read_text(encoding="utf-8")
    tablas = leer_tablas(texto)
    libro = load_workbook(ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.1.xlsx", data_only=True)
    correspondencia = {
        "Decisiones": "D-01", "ADR": "ADR-01", "Supuestos_Economicos": "SE-01",
        "Radar_Data": "T1", "Hipotesis": "H1", "Deuda_Trazabilidad": "DT-01",
        "Dependencias": "F-01", "Fuentes": "FV-01",
    }
    for hoja, primer_id in correspondencia.items():
        esperado = next(t for t in tablas if len(t) > 1 and t[1][0] == primer_id)
        real = [list(f) for f in libro[hoja].values]
        assert real == esperado, f"Diferencias Markdown/Excel: {hoja}"
    por_encabezado = {
        "Canvas_Data": ["Bloque BMC AS-IS"], "Propuestas": ["Propuesta"],
        "Stakeholders": ["Stakeholder"], "Madurez_SocioTecnica": ["Dimensión"],
        "Skills Gap": ["Rol actual del escenario"], "Impact_Canvas": ["Tendencia"],
        "Contribuciones": ["Semana", "Integrante"], "Riesgos": ["Riesgo"],
        "Cambios": ["Cambio"],
    }
    for hoja, encabezado in por_encabezado.items():
        esperado = next(t for t in tablas if t[0][:len(encabezado)] == encabezado)
        assert [list(f) for f in libro[hoja].values] == esperado, f"Diferencias en {hoja}"
    nuevos_supuestos = next(t for t in tablas if len(t) > 1 and t[1][0] == "S-11")
    assert [list(f) for f in list(libro["Supuestos"].values)[11:]] == nuevos_supuestos[1:]
    decisiones = {r[0]: r[2] for r in list(libro["Decisiones"].values)[1:]}
    assert "BMC" in decisiones["D-06"]
    assert "API-first" in decisiones["D-09"]
    assert "BI" in decisiones["D-11"]
    assert "PoC" in decisiones["D-12"]
    assert "Probar" in decisiones["D-13"]
    assert libro["Contribuciones"].max_row == 21
    assert libro["Madurez_SocioTecnica"].max_row == 9
    assert libro["Skills Gap"].max_row == 7
    assert libro["Supuestos"].max_row == 22
    assert len(libro.sheetnames) == 20
    libro.close()

    figuras = []
    for tipo, contenido in bloques(texto):
        if tipo == "imagen":
            alt, relativa = contenido
            ruta = (FUENTE.parent / relativa).resolve()
            figuras.append(ruta)
            assert all(s in alt for s in ("Elaboración propia", "v1.1", "24/10/2026"))
            svg = ruta.with_suffix(".svg")
            ET.fromstring(svg.read_bytes())
            assert "Odoo" not in svg.read_text(encoding="utf-8")
            assert "PostgreSQL" not in svg.read_text(encoding="utf-8")
            with pymupdf.open(ruta.with_suffix(".pdf")) as d:
                assert len(d) == 1
    assert len(figuras) == 4

    docx = DESTINO / (NOMBRE + ".docx")
    with zipfile.ZipFile(docx) as z:
        assert z.testzip() is None
        for nombre in z.namelist():
            if nombre.endswith(".xml"):
                ET.fromstring(z.read(nombre))
    documento = Document(docx)
    assert len(documento.inline_shapes) == 4
    assert len(documento.tables) == len(tablas)
    for original, tabla in zip(tablas, documento.tables):
        assert [[c.text for c in fila.cells] for fila in tabla.rows] == [
            [sin_marcas(c) for c in fila] for fila in original]

    with pymupdf.open(DESTINO / (NOMBRE + ".pdf")) as pdf:
        paginas = len(pdf)
        # Excluir el pie de página antes de unir párrafos partidos entre páginas.
        texto_pdf = "\n".join(p.get_text(clip=pymupdf.Rect(
            0, 0, p.rect.width, p.rect.height - 40)) for p in pdf)
        normal_pdf = normalizar(texto_pdf)
        faltantes = []
        for tipo, valor in bloques(texto):
            fragmentos = ([c for f in valor for c in f] if tipo == "tabla" else
                          [valor[1]] if tipo == "titulo" else
                          [valor[0]] if tipo == "imagen" else [valor])
            for f in fragmentos:
                if normalizar(f) not in normal_pdf:
                    faltantes.append(f)
        assert not faltantes, f"Texto ausente del PDF: {faltantes[:5]}"
        fuera = []
        for n, pagina in enumerate(pdf, 1):
            for x0, y0, x1, y1, *_ in pagina.get_text("blocks"):
                if x0 < 0 or y0 < 0 or x1 > pagina.rect.width + 1 or y1 > pagina.rect.height + 1:
                    fuera.append(n)
        assert not fuera, f"Contenido fuera de página PDF: {fuera}"
        assert sum(len(p.get_images()) for p in pdf) == 4

    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    poc = ROOT / "poc/POC-01-estimacion-historicos"
    comandos = [
        [sys.executable, str(poc / "pruebas/verificar_poc.py")],
        [sys.executable, "-m", "unittest", "discover", "-s", str(poc / "pruebas"), "-p", "test_*.py"],
    ]
    salida = []
    for comando in comandos:
        ejecucion = subprocess.run(comando, cwd=ROOT, env=env, capture_output=True, text=True, encoding="utf-8")
        salida.extend(["Comando: " + " ".join(comando), ejecucion.stdout, ejecucion.stderr])
        assert ejecucion.returncode == 0, "Fallo de PoC/pruebas: " + ejecucion.stderr
    (EVIDENCIAS / "salida_verificacion_poc.txt").write_text("\n".join(salida), encoding="utf-8")
    with (poc / "resultados/resultados.csv").open(encoding="utf-8", newline="") as f:
        filas = list(csv.DictReader(f))
    analitico = sum(float(f["error_modelo"]) for f in filas) / len(filas)
    control = sum(float(f["error_manual"]) for f in filas) / len(filas)
    reduccion = (control - analitico) / control
    assert round(analitico * 100, 2) == 11.31
    assert round(control * 100, 2) == 20.19
    assert round(reduccion * 100, 2) == 43.98
    assert sum(float(f["error_modelo"]) > .15 for f in filas) == 2

    # Ningún archivo previamente versionado de entregas se ha alterado.
    historicos = subprocess.run(["git", "diff", "--name-only", "HEAD", "--", "entregas"],
                               cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    assert not historicos, f"Entrega histórica modificada: {historicos}"

    resumen = {
        "fecha_revision": "2026-10-24", "version": "v1.1", "pdf_paginas": paginas,
        "python": sys.version.split()[0],
        "tablas_sincronizadas_docx": len(tablas), "figuras": 4,
        "hojas_excel": 20, "decisiones": 16, "adr": 6,
        "contribuciones_historicas_declaradas": 20, "pruebas_unitarias": 9,
        "mape_analitico_pct": round(analitico * 100, 2),
        "mape_control_pct": round(control * 100, 2),
        "reduccion_pct": round(reduccion * 100, 2),
        "entregas_historicas_modificadas": False,
        "aprobacion_humana": "pendiente", "h2_h3": "pendientes",
        "limite": "Comprobación técnica; no calificación ni evidencia de impacto empresarial",
    }
    (EVIDENCIAS / "verificacion.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(resumen, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    verificar()
