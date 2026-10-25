#!/usr/bin/env python3
"""Lectura común del Markdown y generación de DOCX/PDF v1.1.

Requiere python-docx, reportlab y Pillow. No convierte ni sobreescribe
entregas históricas. Ambos formatos proceden del mismo Markdown.
"""

from html import escape
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[2]
FUENTE = Path(__file__).with_name("Semana5_Hito_I_v1.1.md")
DESTINO = ROOT / "entregas/semana-05/revision-v1.1"
NOMBRE = "TISG_E01_S05_HitoI_v1.1_2026-10-24"


def bloques(texto):
    lineas = texto.splitlines()
    i = 0
    while i < len(lineas):
        ln = lineas[i].strip()
        if not ln or ln == "---":
            i += 1
            continue
        if ln.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            if any(len(f) != len(filas[0]) for f in filas):
                raise ValueError("Tabla con número de columnas inconsistente")
            yield "tabla", filas
            continue
        imagen = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", ln)
        if imagen:
            yield "imagen", imagen.groups()
        else:
            encabezado = re.match(r"^(#{1,6})\s+(.*)", ln)
            if encabezado:
                yield "titulo", (len(encabezado[1]), encabezado[2])
            else:
                yield "parrafo", ln
        i += 1


def leer_tablas(texto):
    return [contenido for tipo, contenido in bloques(texto) if tipo == "tabla"]


def sin_marcas(texto):
    return texto.replace("**", "").replace("`", "")


def generar_docx(texto):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.8)
    s.left_margin = s.right_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for tipo, valor in bloques(texto):
        if tipo == "titulo":
            nivel, titulo = valor
            doc.add_heading(sin_marcas(titulo), level=min(nivel, 3))
        elif tipo == "tabla":
            tabla = doc.add_table(rows=1, cols=len(valor[0]))
            tabla.style = "Light Shading Accent 1"
            for j, celda in enumerate(valor[0]):
                tabla.rows[0].cells[j].text = sin_marcas(celda)
            repetir = OxmlElement("w:tblHeader")
            tabla.rows[0]._tr.get_or_add_trPr().append(repetir)
            for fila in valor[1:]:
                for celda, contenido in zip(tabla.add_row().cells, fila):
                    celda.text = sin_marcas(contenido)
            for fila in tabla.rows:
                # Una fila no se parte entre páginas.
                fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for celda in fila.cells:
                    for parrafo in celda.paragraphs:
                        for run in parrafo.runs:
                            run.font.size = Pt(8)
            doc.add_paragraph()
        elif tipo == "imagen":
            alt, ruta = valor
            doc.add_picture(str((FUENTE.parent / ruta).resolve()), width=Cm(17.5))
            doc.add_paragraph(alt, style="Caption")
        else:
            p = doc.add_paragraph()
            for fragmento in re.split(r"(\*\*.*?\*\*|`.*?`)", valor):
                run = p.add_run(sin_marcas(fragmento))
                run.bold = fragmento.startswith("**")
    pie = s.footer.paragraphs[0]
    pie.text = "Grupo 1 · Hito I v1.1 · 24/10/2026 · "
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    pie._p.append(campo)
    doc.core_properties.title = "Hito I corregido — Grupo 1 — v1.1"
    doc.core_properties.author = "Grupo 1 — revisión humana pendiente"
    ruta = DESTINO / (NOMBRE + ".docx")
    doc.save(ruta)
    return ruta


def generar_pdf(texto):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Image, Spacer, KeepTogether, PageBreak

    fonts = Path("C:/Windows/Fonts")
    pdfmetrics.registerFont(TTFont("Hito", str(fonts / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("Hito-Bold", str(fonts / "arialbd.ttf")))
    pdfmetrics.registerFontFamily("Hito", normal="Hito", bold="Hito-Bold", italic="Hito", boldItalic="Hito-Bold")
    estilos = getSampleStyleSheet()
    for estilo in estilos.byName.values():
        estilo.fontName = "Hito"
    estilos["BodyText"].fontSize = 9
    estilos["BodyText"].leading = 13
    for nombre in ("Heading1", "Heading2", "Heading3"):
        estilos[nombre].keepWithNext = True
    estilos.add(ParagraphStyle("Celda", fontName="Hito", fontSize=7.5, leading=10, alignment=TA_LEFT,
                              splitLongWords=True))
    estilos.add(ParagraphStyle("PieImagen", fontName="Hito", fontSize=8, leading=10, spaceAfter=10))
    elementos = []
    ancho = A4[0] - 3.4 * cm

    def xml(t):
        t = escape(t).replace("`", "")
        return re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", t)

    contenido = list(bloques(texto))
    for indice, (tipo, valor) in enumerate(contenido):
        if tipo == "titulo":
            if indice + 1 < len(contenido) and contenido[indice + 1][0] == "imagen":
                elementos.append(PageBreak())
            nivel, titulo = valor
            estilo = "Title" if nivel == 1 else "Heading" + str(min(nivel - 1, 3))
            elementos.append(Paragraph(xml(titulo), estilos[estilo]))
        elif tipo == "tabla":
            datos = [[Paragraph(xml(c), estilos["Celda"]) for c in fila] for fila in valor]
            pesos = [max(10, min(52, max(len(sin_marcas(f[j])) for f in valor)))
                     for j in range(len(valor[0]))]
            minimo = 32
            restante = ancho - minimo * len(pesos)
            anchos = [minimo + restante * p / sum(pesos) for p in pesos]
            tabla = LongTable(datos, colWidths=anchos, repeatRows=1, hAlign="LEFT")
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            elementos.extend([tabla, Spacer(1, 8)])
        elif tipo == "imagen":
            from PIL import Image as PILImage
            alt, relativa = valor
            ruta = (FUENTE.parent / relativa).resolve()
            with PILImage.open(ruta) as imagen:
                w, h = imagen.size
            elementos.append(KeepTogether([
                Image(str(ruta), width=ancho, height=ancho * h / w),
                Paragraph(xml(alt), estilos["PieImagen"])]))
        else:
            elementos.append(Paragraph(xml(valor), estilos["BodyText"]))

    def pie(canvas, documento):
        canvas.setFont("Hito", 8)
        canvas.drawString(1.7 * cm, 1.0 * cm, "Grupo 1 · Hito I v1.1 · 24/10/2026 · revisión humana pendiente")
        canvas.drawRightString(A4[0] - 1.7 * cm, 1.0 * cm, str(documento.page))

    ruta = DESTINO / (NOMBRE + ".pdf")
    documento = SimpleDocTemplate(str(ruta), pagesize=A4, leftMargin=1.7 * cm, rightMargin=1.7 * cm,
                                 topMargin=1.7 * cm, bottomMargin=1.7 * cm,
                                 title="Hito I corregido — Grupo 1 — v1.1", author="Grupo 1")
    documento.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return ruta


if __name__ == "__main__":
    DESTINO.mkdir(parents=True, exist_ok=True)
    texto = FUENTE.read_text(encoding="utf-8")
    for resultado in (generar_docx(texto), generar_pdf(texto)):
        print(f"[OK] {resultado.relative_to(ROOT)}")
