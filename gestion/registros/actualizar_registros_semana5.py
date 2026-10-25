#!/usr/bin/env python3
"""Genera SOLO el libro S5 v1.1; conserva los libros históricos.

Requiere openpyxl. Las tablas del Hito I v1.1 son la fuente única de los
registros revisados. No ejecuta el antiguo script S4 ni reasigna IDs.
"""

from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "documentos" / "proyecto"))
from artefactos_hito import leer_tablas  # noqa: E402

DOCUMENTO = ROOT / "documentos/proyecto/Semana5_Hito_I_v1.1.md"
SALIDA = ROOT / "gestion/registros/Registros_y_Radar_Semana5_v1.1.xlsx"


def construir():
    tablas = leer_tablas(DOCUMENTO.read_text(encoding="utf-8"))

    def tabla(inicio):
        candidatas = [t for t in tablas if t[0][0] == inicio]
        if len(candidatas) != 1:
            raise ValueError(f"Se esperaba una tabla {inicio!r}; hay {len(candidatas)}")
        return candidatas[0]

    def id_tabla(identificador):
        candidatas = [t for t in tablas if len(t) > 1 and t[1][0] == identificador]
        if len(candidatas) != 1:
            raise ValueError(f"Tabla no única: {identificador}")
        return candidatas[0]

    libro = Workbook()
    libro.remove(libro.active)

    def agregar(nombre, filas):
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(fila)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        hoja.sheet_view.zoomScale = 85
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E79")
        for fila in hoja.iter_rows():
            for celda in fila:
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        for columna in range(1, hoja.max_column + 1):
            longitud = max(len(str(hoja.cell(f, columna).value or ""))
                           for f in range(1, hoja.max_row + 1))
            hoja.column_dimensions[get_column_letter(columna)].width = min(58, max(15, longitud * .65))
        for fila in range(2, hoja.max_row + 1):
            hoja.row_dimensions[fila].height = 65
        return hoja

    agregar("Control", [
        ["Campo", "Valor"],
        ["Versión", "v1.1 — corrección del Hito I; 24/10/2026"],
        ["Fuente única revisada", str(DOCUMENTO.relative_to(ROOT)).replace("\\", "/")],
        ["Regla", "Históricos intactos; reconstrucción documental, no entrega retroactiva"],
        ["Datos confirmados por equipo", "ERP propio CodeIgniter/MySQL; versiones desconocidas"],
        ["Madurez", "Diagnóstico hipotético, no entrevistas ni medición empresarial"],
        ["Contribuciones", "Declaraciones históricas; no se inventan aportes de la corrección"],
        ["Aprobación", "Revisión y validación humana del equipo pendientes"],
        ["Referencias", "Fuentes consultadas separadas de citas históricas no revalidadas"],
    ])
    agregar("Decisiones", id_tabla("D-01"))
    agregar("ADR", id_tabla("ADR-01"))
    agregar("Radar_Data", id_tabla("T1"))

    historico = load_workbook(ROOT / "gestion/registros/Registros_y_Radar_Semana1.xlsx", data_only=True)
    supuestos = [["ID", "Enunciado", "Tipo", "Estado", "Validación / alcance"]]
    for r in list(historico["Supuestos"].values)[1:]:
        if r[0]:
            supuestos.append([r[0], r[1], r[2], f"Histórico S1: {r[3]}",
                              f"Escenario académico, no nueva verificación. {r[4] or ''} Fuente: {r[5] or ''}"])
    # Corrige referencia interna sin alterar el libro original.
    for fila in supuestos[1:]:
        fila[4] = fila[4].replace("sección 3.3", "sección 3.1")
    supuestos.extend(id_tabla("S-11")[1:])
    agregar("Supuestos", supuestos)
    agregar("Supuestos_Economicos", id_tabla("SE-01"))
    agregar("Contribuciones", next(
        t for t in tablas if t[0][:2] == ["Semana", "Integrante"]))
    agregar("Canvas_Data", tabla("Bloque BMC AS-IS"))
    agregar("Propuestas", tabla("Propuesta"))
    agregar("Stakeholders", tabla("Stakeholder"))
    agregar("Madurez_SocioTecnica", tabla("Dimensión"))
    agregar("Skills Gap", tabla("Rol actual del escenario"))
    agregar("Impact_Canvas", tabla("Tendencia"))
    agregar("Hipotesis", id_tabla("H1"))
    agregar("Deuda_Trazabilidad", id_tabla("DT-01"))
    agregar("Dependencias", id_tabla("F-01"))
    agregar("Riesgos", tabla("Riesgo"))
    agregar("Cambios", tabla("Cambio"))
    agregar("Fuentes", id_tabla("FV-01"))

    fuentes = [["ID original", "Nivel declarado", "Fuente histórica", "Año", "Uso declarado",
                "URL declarada", "Procedencia", "Estado en revisión v1.1"]]
    for nombre in ("Registros_y_Radar_Semana1.xlsx", "Registros_y_Radar_Semana4.xlsx"):
        anterior = load_workbook(ROOT / "gestion/registros" / nombre, data_only=True)
        for r in list(anterior["Fuentes"].values)[1:]:
            if r[0] and not any(f[0] == r[0] for f in fuentes[1:]):
                fuentes.append([*r[:6], nombre, "Cita histórica; no revalidada. No usar cifras sin verificar original"])
        anterior.close()
    agregar("Fuentes_Historicas", fuentes)
    historico.close()

    # Identificadores y número de contribuciones controlados antes de exportar.
    assert [r[0] for r in list(libro["Decisiones"].values)[1:]] == [f"D-{i:02}" for i in range(1, 17)]
    assert libro["Contribuciones"].max_row == 21
    assert libro["Supuestos"].max_row == 22
    libro.save(SALIDA)
    print(f"[OK] {SALIDA.relative_to(ROOT)} ({len(libro.sheetnames)} hojas)")
    return SALIDA


if __name__ == "__main__":
    construir()
